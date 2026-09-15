from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
import copy, sys

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)

# ---- their tokens ----
INK   = RGBColor(0x1B,0x1F,0x3B)
NAVY  = RGBColor(0x1E,0x27,0x61)
BODY_ = RGBColor(0x44,0x50,0x6B)
MUTED = RGBColor(0x87,0x92,0xAC)
CORAL = RGBColor(0xE9,0x4F,0x37)
CARD  = RGBColor(0xF4,0xF6,0xFC)
WASH  = RGBColor(0xFF,0xF3,0xF1)
WHITE = RGBColor(0xFF,0xFF,0xFF)
RULE  = RGBColor(0xD8,0xDE,0xEC)
SERIF, SANS = "Cambria", "Calibri"
W = 13.333

LAYOUT = prs.slides[1].slide_layout
FOOTER = "Vishing Challenge  |  Deepfake Voice Detection for Anti-Vishing Systems"

def tb(s,x,y,w,h,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    b=s.shapes.add_textbox(In(x),In(y),In(w),In(h)); tf=b.text_frame
    tf.word_wrap=True; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=anchor; tf.paragraphs[0].alignment=align
    return tf

def para(tf,first=False,sa=6,line=1.2,align=None):
    p=tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after=Pt(sa); p.line_spacing=line
    if align is not None: p.alignment=align
    return p

def run(p,t,sz=13,b=False,c=BODY_,f=SANS):
    r=p.add_run(); r.text=t
    r.font.size=Pt(sz); r.font.bold=b; r.font.name=f; r.font.color.rgb=c
    return r

def rect(s,x,y,w,h,fill=None,line=None,lw=1.0):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,In(x),In(y),In(w),In(h))
    sh.shadow.inherit=False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.text_frame.text=""
    return sh

def hline(s,x1,y,x2,c=RULE,lw=1.0):
    cn=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,In(x1),In(y),In(x2),In(y))
    cn.line.color.rgb=c; cn.line.width=Pt(lw); return cn

def chrome(s, eyebrow, title_parts):
    tf=tb(s,0.60,0.40,10.0,0.30)
    p=para(tf,first=True,sa=0); run(p,eyebrow,12,True,CORAL,SANS)
    tf=tb(s,0.60,0.68,12.20,0.85)
    p=para(tf,first=True,sa=0,line=1.1)
    for t,c in title_parts: run(p,t,30,True,c,SERIF)
    tf=tb(s,0.50,7.15,9.0,0.30)
    p=para(tf,first=True,sa=0); run(p,FOOTER,9,False,MUTED,SANS)

def notes(s,t):
    s.notes_slide.notes_text_frame.text=t.strip()

# ============================================================ NEW A: the numbers
sA = prs.slides.add_slide(LAYOUT)
chrome(sA,"THE EVIDENCE",[("How far it falls",INK)])

BX0,BX1 = 4.35,12.05
bx = lambda v: BX0+(BX1-BX0)*v/60.0
BH = 0.40

rect(sA,bx(50),1.98,BX1-bx(50),1.98,fill=WASH)
tf=tb(sA,bx(50),1.72,BX1-bx(50),0.24,align=PP_ALIGN.CENTER)
p=para(tf,first=True,sa=0); run(p,"COIN FLIP",9,True,CORAL,SANS)

tf=tb(sA,0.60,1.72,3.6,0.24)
p=para(tf,first=True,sa=0); run(p,"TRAINED ON ASVSPOOF 2019, TESTED ON:",9,True,MUTED,SANS)

for lab,sub,val,col,y in [("its own test set","",1.23,NAVY,2.08),
                          ("another corpus","In-the-Wild",37.15,CORAL,2.74),
                          ("2026 cloned voices","over a real VoIP call",50.28,CORAL,3.40)]:
    tf=tb(sA,0.60,y-0.02,3.5,0.55)
    p=para(tf,first=True,sa=1); run(p,lab,14,True,INK,SANS)
    if sub:
        p=para(tf,sa=0); run(p,sub,10,False,MUTED,SANS)
    rect(sA,BX0,y,max(bx(val)-BX0,0.03),BH,fill=col)
    if val>=45:
        tf=tb(sA,bx(val)-1.35,y+0.06,1.2,0.3,align=PP_ALIGN.RIGHT)
        p=para(tf,first=True,sa=0); run(p,"%.2f%%"%val,17,True,WHITE,SANS)
    else:
        tf=tb(sA,bx(val)+0.16,y+0.05,1.5,0.3)
        p=para(tf,first=True,sa=0); run(p,"%.2f%%"%val,17,True,col,SANS)

hline(sA,0.60,4.16,BX1)
tf=tb(sA,0.60,4.30,3.5,0.55)
p=para(tf,first=True,sa=1); run(p,"same test, retrained",14,True,NAVY,SANS)
p=para(tf,sa=0); run(p,"on modern voices",10,False,MUTED,SANS)
rect(sA,BX0,4.32,bx(5.81)-BX0,BH,fill=NAVY)
tf=tb(sA,bx(5.81)+0.16,4.37,3.6,0.3)
p=para(tf,first=True,sa=0)
run(p,"5.81%",17,True,NAVY,SANS); run(p,"    so it can be fixed",13,False,MUTED,SANS)

hline(sA,BX0,4.92,BX1)
for t in (0,15,30,45,60):
    tf=tb(sA,bx(t)-0.4,4.98,0.8,0.24,align=PP_ALIGN.CENTER)
    p=para(tf,first=True,sa=0); run(p,"%d%%"%t,9,False,MUTED,SANS)
tf=tb(sA,0.60,4.98,3.5,0.24)
p=para(tf,first=True,sa=0); run(p,"equal error rate, lower is better",9,False,MUTED,SANS)

rect(sA,0.60,5.55,12.13,0.95,fill=CARD)
rect(sA,0.60,5.55,0.045,0.95,fill=CORAL)
tf=tb(sA,0.95,5.72,11.5,0.62,anchor=MSO_ANCHOR.MIDDLE)
p=para(tf,first=True,sa=0,line=1.25)
run(p,"A 36-point gap between the benchmark and anything else. ",14,False,BODY_,SANS)
run(p,"Closing that gap is what we are building.",14,True,INK,SANS)

notes(sA,"""
This is the number behind the claim on the previous slide. Walk the bars top to bottom.

Bar 1: RawGAT-ST, 1.23% EER on ASVspoof 2019 LA eval. The published state of the art.
Bar 2: the SAME model, same weights, nothing retrained, tested on In-the-Wild, a corpus of found
celebrity audio. 37.15%. Say "same weights" out loud, that is the point.
Bar 3: an ASVspoof-trained detector against 2026 cloned voices over a real VoIP call. 50.28%,
which is a coin flip. Source is RTCFake, a recent preprint, so say "a recent preprint reports"
rather than presenting it as settled.

Then the navy bar: train the same architecture on modern voices instead and it drops to 5.81%.
So this is not a law of nature, it is a training and data problem. That is the whole reason the
project is worth doing, and it sets up the foundation-model bet two slides later.

If asked whether more data fixes it: measured, it does not. 33.9% to 33.1%. Every ASVspoof split
derives from the same VCTK recordings, so more of it adds no new domain.

EER means equal error rate. 50% is random guessing.
""")

# ============================================================ NEW B: why
sB = prs.slides.add_slide(LAYOUT)
chrome(sB,"DIAGNOSIS",[("It learns the dataset, not the voice",INK)])

tf=tb(sB,0.60,1.62,12.13,0.5)
p=para(tf,first=True,sa=0,line=1.3)
run(p,"Before choosing an architecture we asked what the failing detectors are actually keying on. "
      "The answer explains the collapse, and it tells us what to remove during training.",13,False,BODY_,SANS)

cards=[("SILENCE ALONE","85%","accuracy from a classifier given only the length of the opening silence","Muller et al."),
       ("ATTACK ON SILENCE ONLY","82.2%","attack success from changing only the background noise and the mute segments, never the voice","SiFDetectCracker, ACM MM 2023"),
       ("WHERE IT LOOKS","0.5-0.6 kHz","SHAP analysis: the model attends to the gaps between words and one narrow low band","arXiv:2110.03309")]
cw=(12.13-2*0.35)/3
for i,(tag,big,body,cite) in enumerate(cards):
    x=0.60+i*(cw+0.35)
    rect(sB,x,2.42,cw,2.45,fill=CARD)
    tf=tb(sB,x+0.28,2.66,cw-0.56,1.95)
    p=para(tf,first=True,sa=10); run(p,tag,10,True,MUTED,SANS)
    p=para(tf,sa=12,line=1.0); run(p,big,34,True,CORAL,SANS)
    p=para(tf,sa=10,line=1.3); run(p,body,12.5,False,BODY_,SANS)
    p=para(tf,sa=0); run(p,cite,9,False,MUTED,SANS)

rect(sB,0.60,5.25,12.13,1.05,fill=NAVY)
tf=tb(sB,1.00,5.45,11.4,0.68,anchor=MSO_ANCHOR.MIDDLE)
p=para(tf,first=True,sa=0,line=1.3)
run(p,"The detector wins on what belongs to the corpus, not to the speech. ",14.5,False,RGBColor(0xCA,0xDC,0xFC),SANS)
run(p,"Every one of those cues can be taken away in training.",14.5,True,WHITE,SANS)

notes(sB,"""
Three findings, one conclusion. Do not read the cards, explain each in a sentence.

1. Muller et al. trained a classifier on nothing but the DURATION of the silence at the start of
   the clip. 85% accuracy, 15.1% EER. In ASVspoof, real recordings happen to have longer leading
   silence than fake ones, so silence length leaks the label.
2. SiFDetectCracker, ACM MM 2023. A black-box attack that changes only the background noise and
   the mute segments and never touches the speech itself. 82.2% success against RawNet2,
   RawGAT-ST and others.
3. SHAP analysis, arXiv:2110.03309. The classifier attends to non-speech intervals and one narrow
   band around 0.5 to 0.6 kHz.

Conclusion, say it out loud: a meaningful share of published performance is a dataset artifact,
not voice authenticity. That is precisely why it does not transfer. An attacker can strip the
silence trivially, and a live phone call has no clean studio silence anyway.

This is also why we go to foundation models on the next slide rather than training a detector
from scratch on spectrograms, and why degraded-channel and silence augmentation are in the
engineering plan.
""")

# ============================================================ reorder: A,B go after slide 2
ids = prs.slides._sldIdLst
els = list(ids)
newA, newB = els[-2], els[-1]
ids.remove(newA); ids.remove(newB)
ids.insert(2, newA)
ids.insert(3, newB)

# ============================================================ edits to their slides
def find(sl, needle):
    for sh in sl.shapes:
        if sh.has_text_frame and needle.lower() in sh.text_frame.text.lower():
            return sh
    return None

def retext(sh, parts, size=None, keep=True):
    """Replace a shape's text, reusing the first run's formatting."""
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    proto = p0.runs[0] if p0.runs else None
    nm  = proto.font.name if proto and proto.font.name else SANS
    sz  = size or (proto.font.size.pt if proto and proto.font.size else 13)
    bd  = proto.font.bold if proto else False
    try: col = proto.font.color.rgb
    except Exception: col = BODY_
    for r in list(p0.runs): r._r.getparent().remove(r._r)
    for extra in list(tf.paragraphs)[1:]: extra._p.getparent().remove(extra._p)
    for t,b,c in parts:
        run(p0, t, sz, bd if b is None else b, c or col, nm)

sl = list(prs.slides)

# --- execution-plan slide: it promised to BUILD the system in two days. fix to reality.
ex = sl[10]
h = find(ex,"How the two days break down")
if h: retext(h,[("From proposal to a working detector",None,None)])
plan = [("1","2-DAY SPRINT","Literature & proposal",
         "Survey the generalization failure, fix the evaluation protocol, and commit to an architecture. This deck."),
        ("2","WEEKS 2-4","Baseline",
         "Reproduce published clean-condition EER and validate against the official challenge baseline scores."),
        ("3","WEEKS 5-9","Generalization",
         "Degraded-channel and silence augmentation, foundation-model front-end, leave-one-generator-out. The core work."),
        ("4","WEEKS 10-12","Real-time & write-up",
         "Streaming inference, compression, latency on device, then the report and defence.")]
slots = {}
for sh in ex.shapes:
    if not sh.has_text_frame: continue
    t = sh.text_frame.text.strip()
    for n,(num,head,sub,bodytxt) in enumerate(plan):
        if t == str(n+1): slots.setdefault(n,{})["num"]=sh
        if t.upper().startswith("DAY "+("1" if n<2 else "2")) and "num" in slots.get(n,{}) : pass
for sh in ex.shapes:
    if not sh.has_text_frame: continue
    t = sh.text_frame.text.strip().upper()
    for n,(num,head,sub,bodytxt) in enumerate(plan):
        old_heads = ["DAY 1 — AM","DAY 1 — PM","DAY 2 — AM","DAY 2 — PM"]
        if t == old_heads[n].upper() or t == old_heads[n].replace("—","-").upper():
            retext(sh,[(head,None,None)])
        old_subs = ["Data & baseline","Front-end & head","Real-time + robustness","Integration & rehearsal"]
        if sh.text_frame.text.strip() == old_subs[n]:
            retext(sh,[(sub,None,None)])
        old_bodies = ["Assemble benchmark datasets","Add handcrafted spectral","Streaming inference, compression","Wire the trust score"]
        if sh.text_frame.text.strip().startswith(old_bodies[n]):
            retext(sh,[(bodytxt,None,None)])

# --- evaluation slide: fold the streaming false-alarm point into the metrics card
ev = sl[9]
for sh in ev.shapes:
    if not sh.has_text_frame: continue
    t = sh.text_frame.text.strip()
    if t == "EER & min t-DCF":
        retext(sh,[("EER, min t-DCF, per call-minute",None,None)])
    elif t.startswith("Standard PAD metrics"):
        retext(sh,[("Standard PAD metrics on held-out splits. And a streaming detector is judged "
                    "per call, not per window: at 5% per-window error a 60-second call "
                    "false-alarms 79% of the time.",None,None)])

# --- renumber the page-number boxes
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        if abs(sh.left/914400-12.50)<0.2 and abs(sh.top/914400-7.15)<0.2:
            t=sh.text_frame.text.strip()
            if t.isdigit():
                retext(sh,[(str(i),None,None)])

prs.save(OUT)
print("saved", OUT, len(prs.slides._sldIdLst), "slides")
